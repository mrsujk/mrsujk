from __future__ import annotations

import csv
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
except ImportError:
    DND_FILES = None
    TkinterDnD = None


SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}


def sanitize_filename(name: str) -> str:
    invalid_chars = '<>:"/\\|?*'
    cleaned = "".join("_" if char in invalid_chars else char for char in name)
    cleaned = cleaned.strip().strip(".")
    return cleaned or "arquivo"


def normalize_cell(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n")


def build_output_name(workbook_path: Path, sheet_name: str, total_sheets: int) -> str:
    workbook_name = sanitize_filename(workbook_path.stem)
    if total_sheets == 1:
        return f"{workbook_name}.txt"
    sheet_part = sanitize_filename(sheet_name)
    return f"{workbook_name}__{sheet_part}.txt"


def get_unique_path(target_dir: Path, filename: str) -> Path:
    candidate = target_dir / filename
    index = 1
    while candidate.exists():
        stem = Path(filename).stem
        suffix = Path(filename).suffix
        candidate = target_dir / f"{stem}_{index}{suffix}"
        index += 1
    return candidate


class ExcelToTxtApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("Excel para TXT")
        self.root.geometry("780x520")
        self.root.minsize(680, 460)
        self.root.configure(bg="#f4f1ea")

        self.selected_files: list[Path] = []
        self.drag_and_drop_enabled = TkinterDnD is not None

        self._build_interface()

    def _build_interface(self) -> None:
        title = tk.Label(
            self.root,
            text="Conversor de Excel para TXT",
            font=("Segoe UI", 22, "bold"),
            bg="#f4f1ea",
            fg="#1f2937",
        )
        title.pack(pady=(24, 8))

        subtitle_text = (
            "Carregue um ou mais arquivos Excel e gere arquivos .txt em UTF-8"
        )
        if self.drag_and_drop_enabled:
            subtitle_text += " com suporte a arraste e solte."
        else:
            subtitle_text += "."

        subtitle = tk.Label(
            self.root,
            text=subtitle_text,
            font=("Segoe UI", 10),
            bg="#f4f1ea",
            fg="#4b5563",
        )
        subtitle.pack(pady=(0, 16))

        self.drop_frame = tk.Frame(
            self.root,
            bg="#fffaf0",
            highlightbackground="#c4b79f",
            highlightthickness=2,
            bd=0,
        )
        self.drop_frame.pack(fill="x", padx=30, pady=(0, 16))

        if self.drag_and_drop_enabled:
            self.drop_frame.drop_target_register(DND_FILES)
            self.drop_frame.dnd_bind("<<Drop>>", self.handle_drop)

        self.upload_button = tk.Button(
            self.drop_frame,
            text="Upload ou arraste e solte o arquivo para fazer upload",
            font=("Segoe UI", 12, "bold"),
            bg="#d97706",
            fg="white",
            activebackground="#b45309",
            activeforeground="white",
            relief="flat",
            cursor="hand2",
            padx=18,
            pady=20,
            command=self.select_files,
        )
        self.upload_button.pack(fill="x", padx=18, pady=18)

        if self.drag_and_drop_enabled:
            self.upload_button.drop_target_register(DND_FILES)
            self.upload_button.dnd_bind("<<Drop>>", self.handle_drop)

        self.status_label = tk.Label(
            self.root,
            text="Nenhum arquivo selecionado.",
            font=("Segoe UI", 10, "italic"),
            bg="#f4f1ea",
            fg="#374151",
        )
        self.status_label.pack(anchor="w", padx=32, pady=(0, 10))

        self.file_list = tk.Listbox(
            self.root,
            font=("Consolas", 10),
            height=12,
            activestyle="none",
            bd=0,
            highlightthickness=1,
            highlightbackground="#d1d5db",
        )
        self.file_list.pack(fill="both", expand=True, padx=30, pady=(0, 16))

        self.generate_button = tk.Button(
            self.root,
            text="Gerar",
            font=("Segoe UI", 13, "bold"),
            bg="#2563eb",
            fg="white",
            activebackground="#1d4ed8",
            activeforeground="white",
            relief="flat",
            cursor="hand2",
            padx=28,
            pady=14,
            command=self.generate_txt_files,
        )
        self.generate_button.pack(pady=(0, 26))

    def select_files(self) -> None:
        files = filedialog.askopenfilenames(
            title="Selecione os arquivos Excel",
            filetypes=(
                ("Arquivos Excel", "*.xlsx *.xlsm *.xltx *.xltm"),
                ("Todos os arquivos", "*.*"),
            ),
        )
        if not files:
            return
        self.add_files(Path(file) for file in files)

    def handle_drop(self, event) -> None:
        try:
            dropped_items = self.root.tk.splitlist(event.data)
        except tk.TclError:
            dropped_items = [event.data]
        self.add_files(Path(item) for item in dropped_items)

    def add_files(self, candidates) -> None:
        added = 0
        ignored: list[str] = []

        for file_path in candidates:
            if not file_path.exists() or not file_path.is_file():
                continue
            if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
                ignored.append(file_path.name)
                continue
            if file_path not in self.selected_files:
                self.selected_files.append(file_path)
                added += 1

        self.selected_files.sort(key=lambda path: path.name.lower())
        self.refresh_file_list()

        if ignored:
            messagebox.showwarning(
                "Arquivos ignorados",
                "Os arquivos abaixo nao foram adicionados porque nao sao Excel suportado:\n\n"
                + "\n".join(ignored),
            )

        if added == 0 and not ignored:
            messagebox.showinfo(
                "Sem novidades",
                "Os arquivos selecionados ja estavam na lista.",
            )

    def refresh_file_list(self) -> None:
        self.file_list.delete(0, tk.END)
        if not self.selected_files:
            self.status_label.config(text="Nenhum arquivo selecionado.")
            return

        for file_path in self.selected_files:
            self.file_list.insert(tk.END, str(file_path))

        file_count = len(self.selected_files)
        self.status_label.config(
            text=f"{file_count} arquivo(s) pronto(s) para conversao."
        )

    def generate_txt_files(self) -> None:
        if load_workbook is None:
            messagebox.showerror(
                "Dependencia ausente",
                "A biblioteca openpyxl nao esta instalada.\n\n"
                "Instale as dependencias do requirements.txt antes de usar a conversao.",
            )
            return

        if not self.selected_files:
            messagebox.showwarning(
                "Nenhum arquivo",
                "Selecione pelo menos um arquivo Excel antes de gerar.",
            )
            return

        output_dir = filedialog.askdirectory(title="Escolha a pasta onde salvar")
        if not output_dir:
            return

        output_path = Path(output_dir)
        generated_files: list[Path] = []
        failures: list[str] = []

        try:
            self.generate_button.config(state="disabled", text="Gerando...")
            self.root.update_idletasks()

            for excel_file in self.selected_files:
                try:
                    created = self.convert_workbook(excel_file, output_path)
                    generated_files.extend(created)
                except Exception as exc:
                    failures.append(f"{excel_file.name}: {exc}")
        finally:
            self.generate_button.config(state="normal", text="Gerar")

        summary = [
            f"Arquivos TXT gerados: {len(generated_files)}",
            f"Pasta de destino: {output_path}",
        ]

        if generated_files:
            preview = "\n".join(path.name for path in generated_files[:10])
            if preview:
                summary.append("\nGerados:\n" + preview)
            if len(generated_files) > 10:
                summary.append(f"... e mais {len(generated_files) - 10} arquivo(s).")

        if failures:
            summary.append("\nFalhas:\n" + "\n".join(failures))
            messagebox.showwarning("Conversao finalizada com avisos", "\n".join(summary))
        else:
            messagebox.showinfo("Conversao finalizada", "\n".join(summary))

    def convert_workbook(self, excel_file: Path, output_dir: Path) -> list[Path]:
        workbook = load_workbook(excel_file, data_only=True, read_only=True)
        created_files: list[Path] = []
        worksheets = workbook.worksheets

        for sheet in worksheets:
            filename = build_output_name(excel_file, sheet.title, len(worksheets))
            target_path = get_unique_path(output_dir, filename)

            with target_path.open("w", encoding="utf-8", newline="") as txt_file:
                writer = csv.writer(txt_file, delimiter="\t", lineterminator="\n")
                for row in sheet.iter_rows(values_only=True):
                    writer.writerow([normalize_cell(value) for value in row])

            created_files.append(target_path)

        workbook.close()
        return created_files


def create_root() -> tk.Tk:
    if TkinterDnD is not None:
        return TkinterDnD.Tk()
    return tk.Tk()


def main() -> None:
    root = create_root()
    app = ExcelToTxtApp(root)

    if load_workbook is None:
        messagebox.showwarning(
            "Conversao indisponivel",
            "O aplicativo abriu, mas a leitura de Excel depende da biblioteca openpyxl.\n\n"
            "Instale as dependencias do requirements.txt para habilitar a conversao.",
        )
    elif not app.drag_and_drop_enabled:
        messagebox.showinfo(
            "Arraste e solte opcional",
            "O aplicativo esta funcionando, mas o recurso de arraste e solte depende da biblioteca tkinterdnd2.\n\n"
            "Se quiser esse recurso no .exe, instale as dependencias do requirements.txt antes de gerar.",
        )

    root.mainloop()


if __name__ == "__main__":
    main()
